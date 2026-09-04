"""
Build Postal Code Zones txt from GAR Additional Info (BOC26 Common Rating Table).

Output columns (tab-separated):
  Name | Country | Postal Code | Excluded
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

from build_matrix import (
    DATA_START_ROW,
    DESTINATION_LABEL_COLUMN,
    ORIGIN_CITY_LABEL_COLUMN,
    cell_text,
)
from country_codes import to_country_iso
from location_names import (
    FOREIGN_CITIES_BY_AMBIGUOUS_SUFFIX,
    foreign_country_for_ambiguous_suffix,
    is_boilerplate_cluster_member,
    is_excluded_cluster_member,
    is_misclassified_foreign_us_city,
    location_match_key,
    locations_equivalent,
)
from project_paths import OUTPUT_DIR, RATE_INPUT_DIR, ensure_workspace_dirs
from us_ca_city_states import format_us_ca_postal_city

POSTAL_ZONE_COLUMNS = ("Name", "Country", "Postal Code", "Excluded")
POSTAL_CODE_ZONES_SUFFIX = "_postal_code_zones.txt"
ALL_LANES_PATTERN = re.compile(r"all\s+([A-Za-z]{2})\s+lanes", re.IGNORECASE)
US_ZONE_PATTERN = re.compile(r"^US\s+Zone\s+(.+)$", re.IGNORECASE)
CANADA_ZONE_PATTERN = re.compile(r"^Canada\s+Zone\s+(.+)$", re.IGNORECASE)
EXCEPT_PHRASE = "except for any cities as listed in this common rating table"
JP_LANE_EXCEPTION_PATTERN = re.compile(
    r"except\s+(?:down\s+)?(.+?)(?:\)|$|\.)",
    re.IGNORECASE,
)
METRO_AREA_PATTERN = re.compile(
    r"the\s+metro\s+area\s+and\s+surrounding\s+communities",
    re.IGNORECASE,
)

# States/provinces to append to specific GAR regional zone definitions.
GAR_US_ZONE_STATE_ADDITIONS: dict[str, list[str]] = {
    "2": ["HI"],
}
GAR_CA_ZONE_PROVINCE_ADDITIONS: dict[str, list[str]] = {
    "6": ["NL", "NS", "NB", "PE"],
}

US_STATE_CODES = {
    "AL", "AK", "AZ", "AR", "CA", "CO", "CT", "DE", "DC", "FL", "GA", "HI", "ID", "IL", "IN",
    "IA", "KS", "KY", "LA", "ME", "MD", "MA", "MI", "MN", "MS", "MO", "MT", "NE", "NV", "NH",
    "NJ", "NM", "NY", "NC", "ND", "OH", "OK", "OR", "PA", "RI", "SC", "SD", "TN", "TX", "UT",
    "VT", "VA", "WA", "WV", "WI", "WY",
}
CA_PROVINCE_CODES = {"AB", "BC", "MB", "NB", "NL", "NS", "NT", "NU", "ON", "PE", "QC", "SK", "YT"}

CANADA_NAME_TO_PROVINCE: dict[str, str] = {
    "british columbia": "BC",
    "alberta": "AB",
    "manitoba": "MB",
    "saskatchewan": "SK",
    "ontario": "ON",
    "quebec": "QC",
    "new brunswick": "NB",
    "newfoundland": "NL",
    "nova scotia": "NS",
    "prince edward island": "PE",
    "alaska": "AK",
    "hawaii": "HI",
}

CITY_HIGHLIGHT_FILL = PatternFill("solid", fgColor="DAEEF3")
CITY_HIGHLIGHT_FONT = Font(underline="single", bold=True)

from postal_constants import BULK_COUNTRY_POSTAL_CODE

ALPHANUM_POSTAL = BULK_COUNTRY_POSTAL_CODE

SKIP_ROW_MARKERS = (
    "dimensional weight",
    "rounding rules",
    "inbound to us",
    "outbound from us",
    "international bulk",
    "example",
    "rule",
    "comment",
    "decimal points",
    "applicable for the following",
    "origin taiyuan",
    "destination taiyuan",
    "origin country gb",
)


@dataclass(frozen=True)
class PostalCodeZone:
    name: str
    country: str
    postal_code: str
    excluded: str = ""

    def as_txt_row(self) -> str:
        return "\t".join([self.name, self.country, self.postal_code, self.excluded])


@dataclass
class ParsedCity:
    raw: str
    city: str
    region: str
    country: str


@dataclass
class CityCluster:
    name: str
    country: str
    cities: list[ParsedCity] = field(default_factory=list)


@dataclass
class RegionalZoneDef:
    gar_name: str
    country: str
    postal_code: str
    state_or_prov_codes: set[str]
    uses_table_exclusions: bool


def load_additional_info_raw(file_path: Path) -> pd.DataFrame:
    workbook = pd.ExcelFile(file_path)
    if "Additional Info" not in workbook.sheet_names:
        raise RuntimeError(f"No 'Additional Info' sheet found in {file_path.name}")
    return pd.read_excel(file_path, sheet_name="Additional Info", header=None)


def _normalize_country(value: str) -> str:
    iso = to_country_iso(value)
    if iso:
        return iso.upper()
    text = value.strip().upper()
    return text if len(text) == 2 else text


def _should_skip_row(c0: str, c1: str) -> bool:
    combined = f"{c0} {c1}".lower()
    return any(marker in combined for marker in SKIP_ROW_MARKERS)


def _parse_city_field(text: str, *, default_country: str = "US") -> ParsedCity | None:
    value = cell_text(text)
    if not value:
        return None
    if value.upper() in {"AMR", "EMEIA", "APAC"}:
        return None
    if _should_skip_row(value, ""):
        return None

    parts = [part.strip() for part in value.split(",") if part.strip()]
    if len(parts) >= 2:
        last = parts[-1].upper()
        city_part = ", ".join(parts[:-1]) if len(parts) > 2 else parts[0]
        foreign_country = foreign_country_for_ambiguous_suffix(city_part, last)
        if foreign_country:
            return ParsedCity(raw=value, city=city_part, region="", country=foreign_country)
        if default_country in {"US", "CA"}:
            if last in US_STATE_CODES:
                return ParsedCity(raw=value, city=parts[0], region=last, country="US")
            if last in CA_PROVINCE_CODES:
                return ParsedCity(raw=value, city=parts[0], region=last, country="CA")
        if len(last) == 2:
            if default_country not in {"US", "CA"}:
                resolved = _normalize_country(last)
                if resolved and resolved.upper() != default_country.upper() and last not in US_STATE_CODES:
                    return ParsedCity(
                        raw=value,
                        city=city_part,
                        region="",
                        country=resolved.upper(),
                    )
                return ParsedCity(raw=value, city=value, region="", country=default_country)
            country = _normalize_country(last)
            city = ", ".join(parts[:-1]) if len(parts) > 2 else parts[0]
            return ParsedCity(raw=value, city=city, region="", country=country)

    return ParsedCity(raw=value, city=value, region="", country=default_country)


def _is_regional_zone(c0: str) -> bool:
    return bool(US_ZONE_PATTERN.match(c0) or CANADA_ZONE_PATTERN.match(c0))


def _is_anchor_row(c0: str, c1: str) -> bool:
    if not c0:
        return False
    if c1 and ALL_LANES_PATTERN.search(c1):
        return True
    if "," in c0:
        return True
    return False


def _cluster_name_from_city(city: ParsedCity) -> str:
    text = city.city.strip()
    if text.isupper() and len(text) > 3:
        return text.title()
    return text


def _postal_token(city: ParsedCity, *, common_rating_city: str) -> str:
    if city.country in {"US", "CA"}:
        if city.region:
            city_part = city.city.strip().replace(" ", "")
            return f"{city.region}_{city_part}"
        return format_us_ca_postal_city(city.city, city.country, common_rating_city)
    return city.city.strip()


def _find_table_bounds(raw: pd.DataFrame) -> tuple[int, int]:
    start_row = 0
    for row_idx in range(len(raw)):
        c0 = cell_text(raw.iloc[row_idx, 0])
        if "common rating" in c0.lower():
            start_row = row_idx + 1
            break

    end_row = len(raw)
    for row_idx in range(start_row, len(raw)):
        c0 = cell_text(raw.iloc[row_idx, 0])
        c1 = cell_text(raw.iloc[row_idx, 1]) if raw.shape[1] > 1 else ""
        if c0.lower().startswith("port code") or c0 == "Port Code":
            end_row = row_idx
            break
    return start_row, end_row


def _parse_state_codes_from_description(description: str) -> list[str]:
    text = description.split("-")[0] if "-" in description else description
    text = text.replace(EXCEPT_PHRASE, "")
    codes: list[str] = []
    for token in re.split(r"[,/\s]+", text):
        code = token.strip().upper()
        if code in US_STATE_CODES:
            codes.append(code)
    return codes


def _parse_canada_provinces_from_description(description: str) -> list[str]:
    lower = description.lower()
    if EXCEPT_PHRASE in lower:
        lower = lower.split(EXCEPT_PHRASE)[0]
    codes: list[str] = []
    seen: set[str] = set()
    if "atlantic" in lower:
        for code in ("NL", "NS", "NB", "PE"):
            if code not in seen:
                codes.append(code)
                seen.add(code)
    for name, code in sorted(CANADA_NAME_TO_PROVINCE.items(), key=lambda item: -len(item[0])):
        if name in lower and code not in seen:
            codes.append(code)
            seen.add(code)
    return codes


def _parse_canada_postal_prefix(description: str) -> str:
    codes = _parse_canada_provinces_from_description(description)
    return ", ".join(codes)


def _parse_regional_zone(c0: str, c1: str) -> RegionalZoneDef | None:
    us_match = US_ZONE_PATTERN.match(c0)
    if us_match:
        zone_id = us_match.group(1).strip()
        gar_name = f"GAR-US Zone {zone_id}"
        uses_excl = EXCEPT_PHRASE in c1.lower()
        if zone_id.upper() in {"AK", "HI"}:
            code = zone_id.upper()
            return RegionalZoneDef(
                gar_name=gar_name,
                country="US",
                postal_code=code,
                state_or_prov_codes={code},
                uses_table_exclusions=False,
            )
        state_codes = _parse_state_codes_from_description(c1)
        for extra in GAR_US_ZONE_STATE_ADDITIONS.get(zone_id, []):
            if extra not in state_codes:
                state_codes.append(extra)
        postal = ", ".join(state_codes) if state_codes else c1.split("-")[0].strip().upper()
        return RegionalZoneDef(
            gar_name=gar_name,
            country="US",
            postal_code=postal,
            state_or_prov_codes=set(state_codes),
            uses_table_exclusions=uses_excl,
        )

    ca_match = CANADA_ZONE_PATTERN.match(c0)
    if ca_match:
        zone_id = ca_match.group(1).strip()
        gar_name = f"GAR-CA Zone {zone_id}"
        uses_excl = EXCEPT_PHRASE in c1.lower()
        seen_prov: set[str] = set()
        ordered: list[str] = []
        for code in (
            *_parse_canada_provinces_from_description(c1),
            *GAR_CA_ZONE_PROVINCE_ADDITIONS.get(zone_id, []),
        ):
            if code and code not in seen_prov:
                seen_prov.add(code)
                ordered.append(code)
        prefix = ", ".join(ordered)
        return RegionalZoneDef(
            gar_name=gar_name,
            country="CA",
            postal_code=prefix or c1.strip(),
            state_or_prov_codes=seen_prov,
            uses_table_exclusions=uses_excl,
        )
    return None


def parse_common_rating_table(raw: pd.DataFrame) -> tuple[list[CityCluster], list[RegionalZoneDef], list[ParsedCity]]:
    start_row, end_row = _find_table_bounds(raw)
    clusters: list[CityCluster] = []
    regional: list[RegionalZoneDef] = []
    all_cities: list[ParsedCity] = []

    current: CityCluster | None = None

    def flush_cluster() -> None:
        nonlocal current
        if current and current.cities:
            clusters.append(current)
        current = None

    for row_idx in range(start_row, end_row):
        c0 = cell_text(raw.iloc[row_idx, 0])
        c1 = cell_text(raw.iloc[row_idx, 1]) if raw.shape[1] > 1 else ""
        if not c0 and not c1:
            continue
        if _should_skip_row(c0, c1):
            continue

        if _is_regional_zone(c0):
            flush_cluster()
            zone_def = _parse_regional_zone(c0, c1)
            if zone_def:
                regional.append(zone_def)
            continue

        if _is_anchor_row(c0, c1):
            flush_cluster()
            if c1 and ALL_LANES_PATTERN.search(c1):
                lane_match = ALL_LANES_PATTERN.search(c1)
                anchor = _parse_city_field(c0 or c1)
                if anchor and lane_match:
                    lane_country = _normalize_country(lane_match.group(1))
                    anchor = ParsedCity(
                        raw=anchor.raw,
                        city=anchor.city,
                        region=anchor.region,
                        country=lane_country,
                    )
                    all_cities.append(anchor)
                continue
            anchor = _parse_city_field(c0 or c1)
            if anchor is None:
                continue
            cluster_name = _cluster_name_from_city(anchor)
            current = CityCluster(name=cluster_name, country=anchor.country)
            current.cities.append(anchor)
            all_cities.append(anchor)
            if c1 and c1 != c0:
                extra = _parse_city_field(c1, default_country=anchor.country)
                if extra:
                    current.cities.append(extra)
                    all_cities.append(extra)
            continue

        if current is None:
            lone = _parse_city_field(c0 or c1)
            if lone:
                current = CityCluster(name=_cluster_name_from_city(lone), country=lone.country)
                current.cities.append(lone)
                all_cities.append(lone)
            continue

        for value in (c0, c1):
            if not value:
                continue
            member = _parse_city_field(value, default_country=current.country)
            if member is None:
                continue
            if any(location_match_key(m.raw) == location_match_key(member.raw) for m in current.cities):
                continue
            current.cities.append(member)
            all_cities.append(member)

    flush_cluster()
    return clusters, regional, all_cities


def _cluster_to_zone(cluster: CityCluster) -> PostalCodeZone:
    tokens: list[str] = []
    seen: set[str] = set()
    for city in cluster.cities:
        if is_excluded_cluster_member(cluster.name, city.raw):
            continue
        token = _postal_token(city, common_rating_city=cluster.name)
        key = token.casefold()
        if key in seen:
            continue
        seen.add(key)
        tokens.append(token)
    return PostalCodeZone(
        name=cluster.name,
        country=cluster.country,
        postal_code=", ".join(tokens),
        excluded="",
    )


def _jp_all_lanes_exclusions(description: str) -> list[str]:
    """Parse cities/regions excluded from an All JP lanes row (e.g. Except Osaka)."""
    if not description or "except" not in description.lower():
        return []
    if "jp" not in description.lower() and not ALL_LANES_PATTERN.search(description):
        return []
    match = JP_LANE_EXCEPTION_PATTERN.search(description)
    if not match:
        return []
    fragment = match.group(1).strip(" )")
    known: dict[str, str] = {
        "osaka": "Osaka",
        "central": "Central",
    }
    exclusions: list[str] = []
    lowered = fragment.casefold()
    for key, label in known.items():
        if key in lowered:
            exclusions.append(label)
    if exclusions:
        return exclusions
    return [fragment] if fragment else []


def _jp_all_lanes_zone(name: str, *, lane_description: str) -> PostalCodeZone | None:
    exclusions = _jp_all_lanes_exclusions(lane_description)
    if not exclusions:
        return None
    return PostalCodeZone(
        name=name,
        country="JP",
        postal_code=ALPHANUM_POSTAL,
        excluded=", ".join(exclusions),
    )


def _jp_excluded_city_zone(city_name: str) -> PostalCodeZone:
    return PostalCodeZone(
        name=city_name,
        country="JP",
        postal_code=city_name,
        excluded="",
    )


# Cities excluded from an All JP lanes row that also need their own postal zone.
JP_EXCLUDED_COMPANION_CITIES = frozenset({"Osaka"})


def _jp_exclusion_companion_zones(
    exclusions: list[str],
    existing_zones: list[PostalCodeZone],
) -> list[PostalCodeZone]:
    existing_keys = {(location_match_key(zone.name), zone.country.upper()) for zone in existing_zones}
    companions: list[PostalCodeZone] = []
    for city in exclusions:
        if city not in JP_EXCLUDED_COMPANION_CITIES:
            continue
        key = (location_match_key(city), "JP")
        if key in existing_keys:
            continue
        companions.append(_jp_excluded_city_zone(city))
        existing_keys.add(key)
    return companions


def _all_lanes_zone(
    anchor: ParsedCity,
    *,
    lane_country: str | None = None,
    lane_description: str = "",
) -> PostalCodeZone:
    country = lane_country or anchor.country
    name = _cluster_name_from_city(anchor)
    if country == "JP":
        jp_zone = _jp_all_lanes_zone(name, lane_description=lane_description)
        if jp_zone is not None:
            return jp_zone
    return _country_postal_zone(name, country)


def _country_postal_zone(name: str, country: str) -> PostalCodeZone:
    return PostalCodeZone(
        name=name,
        country=country.upper(),
        postal_code=ALPHANUM_POSTAL,
        excluded="-",
    )


@dataclass(frozen=True)
class CountryReference:
  by_name: dict[str, str]
  display_by_iso: dict[str, str]


def _is_iso_country_code(text: str) -> bool:
    return len(text) == 2 and text.isalpha()


def _parse_country_name_code(name: str, code: str) -> tuple[str, str] | None:
    name = cell_text(name)
    code = cell_text(code)
    if not name or not _is_iso_country_code(code):
        return None
    lower_name = name.lower()
    if lower_name in {"country", "iso code", "world countries and iso codes"}:
        return None
    if "," in name:
        return None
    if name.upper() in US_STATE_CODES:
        return None

    iso = code.upper()
    return name, iso


def extract_country_reference(raw: pd.DataFrame) -> CountryReference:
    """Country names and ISO codes from Additional Info (columns B/C and C/D)."""
    by_name: dict[str, str] = {}
    display_by_iso: dict[str, str] = {}
    start_row, end_row = _find_table_bounds(raw)
    for row_idx in range(start_row, end_row):
        cells = [cell_text(raw.iloc[row_idx, col_idx]) for col_idx in range(raw.shape[1])]
        for col_idx in range(len(cells) - 1):
            parsed = _parse_country_name_code(cells[col_idx], cells[col_idx + 1])
            if parsed is None:
                continue
            display, iso = parsed
            by_name[display.casefold()] = iso
            display_by_iso[iso] = display

    return CountryReference(by_name=by_name, display_by_iso=display_by_iso)


def _label_represents_country(label: str, country: str, country_ref: CountryReference) -> bool:
    if not label or not country or len(country) != 2:
        return False
    country_iso = country.upper()
    key = label.casefold()
    if country_ref.by_name.get(key) == country_iso:
        return True
    resolved = to_country_iso(label)
    if len(resolved) == 2 and resolved.upper() == country_iso:
        return True
    if key == country_iso.casefold() and len(label) == 2:
        return True
    return False


def _zones_from_country_reference(country_ref: CountryReference) -> list[PostalCodeZone]:
    zones: list[PostalCodeZone] = []
    for iso, display in country_ref.display_by_iso.items():
        zones.append(_country_postal_zone(display, iso))
    return zones


def _cluster_regions(cluster: CityCluster) -> set[str]:
    regions = {city.region for city in cluster.cities if city.region}
    if regions:
        return regions
    for city in cluster.cities:
        token = _postal_token(city, common_rating_city=cluster.name)
        if "_" in token:
            regions.add(token.split("_", 1)[0])
    return regions


def _build_regional_zones(
    regional_defs: list[RegionalZoneDef],
    clusters: list[CityCluster],
) -> list[PostalCodeZone]:
    zones: list[PostalCodeZone] = []
    for zone_def in regional_defs:
        excluded_tokens: list[str] = []
        seen_excl: set[str] = set()
        if zone_def.state_or_prov_codes:
            for cluster in clusters:
                if cluster.country != zone_def.country:
                    continue
                if not _cluster_regions(cluster) & zone_def.state_or_prov_codes:
                    continue
                for city in cluster.cities:
                    if is_misclassified_foreign_us_city(city.city, city.region, city.country):
                        continue
                    if is_excluded_cluster_member(cluster.name, city.raw):
                        continue
                    token = _postal_token(city, common_rating_city=cluster.name)
                    key = token.casefold()
                    if key in seen_excl:
                        continue
                    seen_excl.add(key)
                    excluded_tokens.append(token)

        zones.append(
            PostalCodeZone(
                name=zone_def.gar_name,
                country=zone_def.country,
                postal_code=zone_def.postal_code,
                excluded=", ".join(excluded_tokens),
            )
        )
    return zones


def build_postal_code_zones_from_raw(raw: pd.DataFrame) -> list[PostalCodeZone]:
    clusters, regional_defs, all_cities = parse_common_rating_table(raw)

    zones: list[PostalCodeZone] = []

    for cluster in clusters:
        zones.append(_cluster_to_zone(cluster))

    # Type C rows: All XX lanes anchors
    start_row, end_row = _find_table_bounds(raw)
    for row_idx in range(start_row, end_row):
        c0 = cell_text(raw.iloc[row_idx, 0])
        c1 = cell_text(raw.iloc[row_idx, 1]) if raw.shape[1] > 1 else ""
        if not c0 or not c1:
            continue
        match = ALL_LANES_PATTERN.search(c1)
        if not match:
            continue
        anchor = _parse_city_field(c0)
        if anchor is None:
            continue
        lane_country = _normalize_country(match.group(1))
        zones.append(_all_lanes_zone(anchor, lane_country=lane_country, lane_description=c1))

    zones.extend(_build_regional_zones(regional_defs, clusters))

    country_ref = extract_country_reference(raw)
    zones.extend(_zones_from_country_reference(country_ref))

    zones = _apply_special_zone_overrides(zones, raw)

    jp_exclusions: list[str] = []
    for row_idx in range(start_row, end_row):
        c1 = cell_text(raw.iloc[row_idx, 1]) if raw.shape[1] > 1 else ""
        lane_match = ALL_LANES_PATTERN.search(c1)
        if not lane_match:
            continue
        if _normalize_country(lane_match.group(1)) == "JP":
            jp_exclusions.extend(_jp_all_lanes_exclusions(c1))
    zones.extend(_jp_exclusion_companion_zones(jp_exclusions, zones))

    zones.sort(key=lambda zone: (zone.country, zone.name.casefold()))
    return zones


def _apply_special_zone_overrides(zones: list[PostalCodeZone], raw: pd.DataFrame) -> list[PostalCodeZone]:
    start_row, end_row = _find_table_bounds(raw)
    tokyo_lane_description = ""
    for row_idx in range(start_row, end_row):
        c0 = cell_text(raw.iloc[row_idx, 0])
        c1 = cell_text(raw.iloc[row_idx, 1]) if raw.shape[1] > 1 else ""
        if location_match_key(c0) == "tokyo" and ALL_LANES_PATTERN.search(c1):
            tokyo_lane_description = c1
            break

    updated: list[PostalCodeZone] = []
    for zone in zones:
        if (
            tokyo_lane_description
            and location_match_key(zone.name) == "tokyo"
            and zone.country == "JP"
        ):
            jp_zone = _jp_all_lanes_zone(zone.name, lane_description=tokyo_lane_description)
            if jp_zone is not None:
                updated.append(jp_zone)
                continue
        if (
            location_match_key(zone.name)
            in FOREIGN_CITIES_BY_AMBIGUOUS_SUFFIX.get("IN", frozenset())
            and zone.country == "US"
        ):
            cleaned_tokens = []
            for part in zone.postal_code.split(", "):
                token = part.strip()
                if token.upper().startswith("IN_"):
                    token = token[3:]
                cleaned_tokens.append(token)
            updated.append(
                PostalCodeZone(
                    name=zone.name,
                    country="IN",
                    postal_code=", ".join(cleaned_tokens),
                    excluded=zone.excluded,
                )
            )
            continue
        updated.append(zone)
    return updated


def _matrix_city_labels(
    matrix_df: pd.DataFrame,
    *,
    label_column: str,
    country_column: str,
) -> list[tuple[str, str]]:
    seen: set[str] = set()
    results: list[tuple[str, str]] = []
    for _, row in matrix_df.iterrows():
        label = cell_text(row.get(label_column))
        country = _normalize_country(cell_text(row.get(country_column)))
        if not label or not country:
            continue
        key = location_match_key(label)
        if key in seen:
            continue
        seen.add(key)
        results.append((label, country))
    return results


def matrix_destination_labels(matrix_df: pd.DataFrame) -> list[tuple[str, str]]:
    return _matrix_city_labels(
        matrix_df,
        label_column="Destination Label",
        country_column="Destination Country",
    )


def matrix_origin_labels(matrix_df: pd.DataFrame) -> list[tuple[str, str]]:
    return _matrix_city_labels(
        matrix_df,
        label_column="Origin City Label",
        country_column="Origin Country",
    )


def _is_valid_zone_name(name: str) -> bool:
    text = cell_text(name)
    if not text or len(text) < 2:
        return False
    lower = text.casefold()
    if re.match(r"^\d+\)", lower):
        return False
    if any(marker in lower for marker in SKIP_ROW_MARKERS):
        return False
    if "origin " in lower and "destination" in lower:
        return False
    if METRO_AREA_PATTERN.search(text):
        return False
    if is_boilerplate_cluster_member(text):
        return False
    if len(text) > 100:
        return False
    return True


def filter_valid_postal_zones(zones: list[PostalCodeZone]) -> list[PostalCodeZone]:
    seen: set[tuple[str, str]] = set()
    kept: list[PostalCodeZone] = []
    for zone in zones:
        if not _is_valid_zone_name(zone.name):
            continue
        key = (location_match_key(zone.name), zone.country)
        if key in seen:
            continue
        seen.add(key)
        kept.append(zone)
    kept.sort(key=lambda z: (z.country, z.name.casefold()))
    return kept


def _lookup_zone(label: str, country: str, zones_by_name: dict[str, PostalCodeZone]) -> PostalCodeZone | None:
    key = location_match_key(label)
    if key in zones_by_name:
        return zones_by_name[key]
    for zone in zones_by_name.values():
        if zone.country == country and location_match_key(zone.name) == key:
            return zone
    for zone in zones_by_name.values():
        if zone.country == country and locations_equivalent(label, zone.name):
            return zone
    return None


def build_final_postal_zones(
    raw_additional_info: pd.DataFrame,
    matrix_df: pd.DataFrame,
    *,
    include_origin_city: bool = True,
) -> tuple[list[PostalCodeZone], list[PostalCodeZone]]:
    all_zones = filter_valid_postal_zones(build_postal_code_zones_from_raw(raw_additional_info))
    zones_by_name: dict[str, PostalCodeZone] = {}
    for zone in all_zones:
        zones_by_name[location_match_key(zone.name)] = zone
    country_ref = extract_country_reference(raw_additional_info)

    final_zones: list[PostalCodeZone] = []

    def append_for_labels(labels: list[tuple[str, str]]) -> None:
        for label, country in labels:
            zone = _lookup_zone(label, country, zones_by_name)
            if zone is not None:
                final_zones.append(zone)
                continue
            if _label_represents_country(label, country, country_ref):
                final_zones.append(_country_postal_zone(label, country))
                continue
            if country in {"US", "CA"}:
                postal = format_us_ca_postal_city(label, country, label)
            else:
                postal = label
            final_zones.append(PostalCodeZone(name=label, country=country, postal_code=postal, excluded=""))

    if include_origin_city:
        append_for_labels(matrix_origin_labels(matrix_df))
    append_for_labels(matrix_destination_labels(matrix_df))

    deduped: list[PostalCodeZone] = []
    seen_keys: set[tuple[str, str]] = set()
    for zone in final_zones:
        key = (location_match_key(zone.name), zone.country)
        if key in seen_keys:
            continue
        seen_keys.add(key)
        deduped.append(zone)

    deduped.sort(key=lambda zone: (zone.country, zone.name.casefold()))
    return deduped, all_zones


def load_business_rules_postal_zones(reference_workbook: Path) -> list[PostalCodeZone]:
    """Postal zones from Rate Card Export **Business rules** sheet (RAM parity)."""
    workbook = pd.ExcelFile(reference_workbook)
    if "Business rules" not in workbook.sheet_names:
        raise RuntimeError(f"No 'Business rules' sheet in {reference_workbook.name}")
    df = pd.read_excel(reference_workbook, sheet_name="Business rules", header=5)
    zones: list[PostalCodeZone] = []
    for _, row in df.iterrows():
        name = cell_text(row.get("Name"))
        country = _normalize_country(cell_text(row.get("Country")))
        postal = cell_text(row.get("Postal Code"))
        if not name or not country:
            continue
        excluded = cell_text(row.get("Excluded"))
        zones.append(
            PostalCodeZone(
                name=name,
                country=country,
                postal_code=postal,
                excluded=excluded,
            )
        )
    return filter_valid_postal_zones(zones)


def compare_postal_zones_to_business_rules(
    generated_zones: list[PostalCodeZone],
    reference_workbook: Path,
) -> dict[str, int | list[str]]:
    """Compare pipeline txt zones to Business rules sheet (same role as RAM export)."""
    reference_zones = load_business_rules_postal_zones(reference_workbook)

    def zone_key(zone: PostalCodeZone) -> tuple[str, str]:
        return (location_match_key(zone.name), zone.country.upper())

    gen_map = {zone_key(z): z for z in generated_zones}
    ref_map = {zone_key(z): z for z in reference_zones}
    only_ref = sorted(ref_map.keys() - gen_map.keys())
    only_gen = sorted(gen_map.keys() - ref_map.keys())
    matched = ref_map.keys() & gen_map.keys()
    postal_mismatch = 0
    for key in matched:
        if gen_map[key].postal_code.strip() != ref_map[key].postal_code.strip():
            postal_mismatch += 1
    return {
        "reference_count": len(reference_zones),
        "generated_count": len(generated_zones),
        "matched_names": len(matched),
        "only_in_reference": len(only_ref),
        "only_in_generated": len(only_gen),
        "postal_code_mismatches": postal_mismatch,
        "sample_only_reference": [ref_map[k].name for k in only_ref[:10]],
        "sample_only_generated": [gen_map[k].name for k in only_gen[:10]],
    }


def write_postal_code_zones_txt(zones: list[PostalCodeZone], output_path: Path) -> Path:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    lines = ["\t".join(POSTAL_ZONE_COLUMNS)]
    lines.extend(zone.as_txt_row() for zone in zones)
    output_path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return output_path


def default_postal_zones_path(matrix_path: Path) -> Path:
    stem = matrix_path.stem.replace("_matrix", "")
    return OUTPUT_DIR / f"{stem}{POSTAL_CODE_ZONES_SUFFIX}"


def _city_column_index(column_name: str) -> int:
    from build_matrix import EXPORT_HEADERS, MATRIX_COLUMNS

    worksheet_headers = [EXPORT_HEADERS.get(header, header) for header in MATRIX_COLUMNS]
    return worksheet_headers.index(column_name) + 1


def highlight_labeled_city_column(
    worksheet,
    *,
    column_name: str,
    zone_names_set: set[str],
) -> int:
    col_index = _city_column_index(column_name)
    highlighted = 0
    names_cf = {name.casefold() for name in zone_names_set}

    for row_index in range(DATA_START_ROW, worksheet.max_row + 1):
        cell = worksheet.cell(row_index, col_index)
        value = cell_text(cell.value)
        if not value or value.casefold() not in names_cf:
            continue
        cell.fill = CITY_HIGHLIGHT_FILL
        cell.font = CITY_HIGHLIGHT_FONT
        highlighted += 1
    return highlighted


def highlight_city_labels_in_matrix(
    matrix_path: Path,
    *,
    zone_names: set[str],
) -> tuple[int, int]:
    workbook = load_workbook(matrix_path)
    worksheet = workbook["Rate card"]
    origin = highlight_labeled_city_column(
        worksheet,
        column_name=ORIGIN_CITY_LABEL_COLUMN,
        zone_names_set=zone_names,
    )
    destination = highlight_labeled_city_column(
        worksheet,
        column_name=DESTINATION_LABEL_COLUMN,
        zone_names_set=zone_names,
    )
    workbook.save(matrix_path)
    return origin, destination


def resolve_rate_workbook(
    processing_path: Path,
    rate_file_path: Path | None = None,
) -> Path:
    if rate_file_path is not None and rate_file_path.is_file():
        return rate_file_path

    stem = processing_path.stem.replace("_extracted", "")
    candidates = sorted(RATE_INPUT_DIR.glob("*.xlsx"), key=lambda p: p.stat().st_mtime, reverse=True)
    for path in candidates:
        if path.name.startswith("~$"):
            continue
        if path.stem == stem or stem.startswith(path.stem) or path.stem.startswith(stem.split()[0]):
            return path

    for path in candidates:
        if path.name.startswith("~$"):
            continue
        if stem.split()[0] in path.name:
            return path

    raise FileNotFoundError(
        f"Could not find rate workbook for processing file {processing_path.name} in {RATE_INPUT_DIR}"
    )


def run_build_postal_code_zones(
    *,
    rate_file_path: Path,
    matrix_path: Path,
    matrix_df: pd.DataFrame,
    output_path: Path | None = None,
    include_origin_city: bool = True,
) -> Path:
    ensure_workspace_dirs()
    raw = load_additional_info_raw(rate_file_path)
    matrix_matched, catalog_zones = build_final_postal_zones(
        raw,
        matrix_df,
        include_origin_city=include_origin_city,
    )
    export_zones = catalog_zones
    txt_path = output_path or default_postal_zones_path(matrix_path)
    write_postal_code_zones_txt(export_zones, txt_path)

    zone_names = {zone.name for zone in matrix_matched}
    origin_h, dest_h = highlight_city_labels_in_matrix(matrix_path, zone_names=zone_names)

    print(f"  Wrote {len(export_zones)} postal code zone(s) (BOC26 catalog; {len(matrix_matched)} tied to matrix labels)")

    ref_candidates = sorted(
        matrix_path.parent.parent.glob("Rate Card Export*.xlsx"),
        key=lambda p: p.stat().st_mtime,
        reverse=True,
    )
    if ref_candidates:
        try:
            stats = compare_postal_zones_to_business_rules(export_zones, ref_candidates[0])
            print(
                f"  vs Business rules ({ref_candidates[0].name}): "
                f"matched {stats['matched_names']}/{stats['reference_count']}, "
                f"only ref {stats['only_in_reference']}, only txt {stats['only_in_generated']}, "
                f"postal mismatches {stats['postal_code_mismatches']}"
            )
        except Exception as exc:
            print(f"  (Business rules compare skipped: {exc})")
    print(f"  Highlighted {origin_h} origin and {dest_h} destination label cell(s) in matrix")
    print(f"  Saved postal code zones to: {txt_path}")
    return txt_path
