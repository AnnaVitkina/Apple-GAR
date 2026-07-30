"""Highlight duplicate shipment lanes in the pipeline matrix workbook."""

from __future__ import annotations

from collections import defaultdict
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

from build_matrix import (
    DATA_START_ROW,
    MATRIX_COLUMNS,
    cell_text,
)

DUPLICATE_LANE_FILL = PatternFill("solid", fgColor="C6E0B4")
LEGEND_FONT = Font(bold=True, size=10)

SHIPMENT_IDENTITY_FIELDS: tuple[str, ...] = (
    "Tab",
    "Origin Country",
    "Origin State",
    "Origin City",
    "Origin City Label",
    "Uplift Airport",
    "Origin Airport",
    "Rate Type",
    "Destination Country",
    "Destination State",
    "Destination",
    "Destination Label",
    "Destination Airport",
    "Country of Clearance",
    "Service",
    "Lane Type",
    "Service Product",
    "Measurement Type",
    "Business Segment",
    "Invoice type",
    "Valid to",
)


def _normalize_field_value(value: object) -> str:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ""
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return f"{float(value):.4f}".rstrip("0").rstrip(".")
    return cell_text(value)


def _shipment_identity_key(row: pd.Series) -> tuple[str, ...]:
    return tuple(
        _normalize_field_value(row.get(field)) for field in SHIPMENT_IDENTITY_FIELDS
    )


def _duplicate_lane_columns() -> list[int]:
    columns = [MATRIX_COLUMNS.index("Lane #") + 1]
    for field in SHIPMENT_IDENTITY_FIELDS:
        if field in MATRIX_COLUMNS:
            columns.append(MATRIX_COLUMNS.index(field) + 1)
    return columns


def _highlight_duplicate_lanes(worksheet, matrix_df: pd.DataFrame) -> dict[str, int]:
    buckets: dict[tuple[str, ...], list[int]] = defaultdict(list)
    for offset in range(len(matrix_df)):
        buckets[_shipment_identity_key(matrix_df.iloc[offset])].append(offset)

    offsets: set[int] = set()
    group_count = 0
    for group in buckets.values():
        if len(group) > 1:
            group_count += 1
            offsets.update(group)

    columns = _duplicate_lane_columns()
    for offset in offsets:
        excel_row = DATA_START_ROW + offset
        for col_index in columns:
            worksheet.cell(excel_row, col_index).fill = DUPLICATE_LANE_FILL

    return {
        "duplicate_rows": len(offsets),
        "duplicate_groups": group_count,
    }


def _write_legend(worksheet, stats: dict[str, int]) -> None:
    worksheet.cell(4, 1, "Green = duplicate shipment details (same values as another row in this file).")
    worksheet.cell(4, 1).font = LEGEND_FONT
    worksheet.cell(
        5,
        1,
        f"Duplicate rows: {stats.get('duplicate_rows', 0)} "
        f"({stats.get('duplicate_groups', 0)} groups)",
    )


def apply_matrix_highlights(matrix_path: Path, matrix_df: pd.DataFrame) -> dict[str, int]:
    workbook = load_workbook(matrix_path)
    worksheet = workbook["Rate card"]
    stats = _highlight_duplicate_lanes(worksheet, matrix_df)
    _write_legend(worksheet, stats)
    workbook.save(matrix_path)
    return stats


def apply_matrix_highlights_if_saved(
    matrix_path: Path,
    matrix_df: pd.DataFrame,
) -> dict[str, int]:
    stats = apply_matrix_highlights(matrix_path, matrix_df)
    print(
        f"  Matrix highlight: duplicate rows {stats.get('duplicate_rows', 0)} "
        f"({stats.get('duplicate_groups', 0)} groups)"
    )
    return stats
